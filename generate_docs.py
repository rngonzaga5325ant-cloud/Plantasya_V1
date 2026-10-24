import os
import re
import html
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def md_to_reportlab_html(text):
    text = html.escape(text)
    # Convert bold **text** to <b>text</b>
    text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', text)
    # Convert italic *text* to <i>text</i>
    text = re.sub(r'\*(.*?)\*', r'<i>\1</i>', text)
    # Convert inline code `text` to font tag
    text = re.sub(r'`(.*?)`', r'<font face="Courier" color="#1b4332">\1</font>', text)
    return text

def md_to_docx_text(text):
    text = text.replace("**", "").replace("*", "").replace("`", "")
    return text

def create_pdf(filename):
    doc = SimpleDocTemplate(
        filename,
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    primary_color = colors.HexColor("#1b4332")   # Dark Forest Green
    secondary_color = colors.HexColor("#2d6a4f") # Medium Forest Green
    accent_color = colors.HexColor("#40916c")    # Light Green
    text_dark = colors.HexColor("#212529")       # Near Black
    bg_light = colors.HexColor("#f8f9fa")        # Off-white
    border_color = colors.HexColor("#d8f3dc")    # Light Green border

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=22,
        leading=26,
        textColor=primary_color,
        alignment=1, # Center
        spaceAfter=10
    )

    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=secondary_color,
        alignment=1, # Center
        spaceAfter=15
    )

    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#495057"),
        alignment=1,
        spaceAfter=20
    )

    h1_style = ParagraphStyle(
        'Header1',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=primary_color,
        spaceBefore=16,
        spaceAfter=8,
        keepWithNext=True
    )

    h2_style = ParagraphStyle(
        'Header2',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11.5,
        leading=15,
        textColor=secondary_color,
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'BodyTextCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=text_dark,
        spaceAfter=6
    )

    bullet_style = ParagraphStyle(
        'BulletCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=text_dark,
        leftIndent=15,
        spaceAfter=4
    )

    code_style = ParagraphStyle(
        'CodeCustom',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=8.5,
        leading=11,
        textColor=primary_color,
        backColor=bg_light,
        borderColor=border_color,
        borderWidth=1,
        borderPadding=5,
        spaceAfter=6
    )

    story = []

    story.append(Paragraph("PLANTASYA MOBILE APPLICATION", title_style))
    story.append(Paragraph("OFFICIAL COMPREHENSIVE USER MANUAL & OPERATING GUIDE", subtitle_style))
    story.append(Paragraph("<b>App Name:</b> Plantasya | <b>Version:</b> 1.0.0 | <b>Target OS:</b> Android 8.0+<br/><b>Package:</b> com.example.plantasya_mobileapp", meta_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=primary_color, spaceBefore=0, spaceAfter=15))

    with open("PLANTASYA_USER_MANUAL.md", "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_table = False
    table_data = []

    for line in lines:
        line_str = line.strip()

        if line_str.startswith("# "):
            if in_table:
                t = Table(table_data, colWidths=[140, 150, 214])
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), primary_color),
                    ('TEXTCOLOR', (0,0), (-1,0), colors.white),
                    ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0,0), (-1,-1), 8.5),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 5),
                    ('TOPPADDING', (0,0), (-1,-1), 5),
                    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#ced4da")),
                    ('BACKGROUND', (0,1), (-1,-1), colors.white)
                ]))
                story.append(t)
                story.append(Spacer(1, 8))
                in_table = False
                table_data = []

            story.append(Paragraph(md_to_reportlab_html(line_str[2:]), h1_style))
            story.append(HRFlowable(width="100%", thickness=1, color=accent_color, spaceBefore=1, spaceAfter=8))

        elif line_str.startswith("## "):
            story.append(Paragraph(md_to_reportlab_html(line_str[3:]), h2_style))

        elif line_str.startswith("### "):
            story.append(Paragraph(md_to_reportlab_html(line_str[4:]), h2_style))

        elif line_str.startswith("|"):
            if "---" in line_str:
                continue
            cells = [c.strip() for c in line_str.split("|")[1:-1]]
            if len(cells) > 0:
                in_table = True
                p_cells = [Paragraph(md_to_reportlab_html(c), body_style) for c in cells]
                table_data.append(p_cells)

        elif line_str.startswith("* ") or line_str.startswith("- "):
            if in_table:
                t = Table(table_data, colWidths=[140, 150, 214])
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), primary_color),
                    ('TEXTCOLOR', (0,0), (-1,0), colors.white),
                    ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0,0), (-1,-1), 8.5),
                    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#ced4da")),
                ]))
                story.append(t)
                story.append(Spacer(1, 8))
                in_table = False
                table_data = []

            bullet_text = md_to_reportlab_html(line_str[2:])
            story.append(Paragraph(f"• {bullet_text}", bullet_style))

        elif re.match(r'^\d+\.\s', line_str):
            if in_table:
                t = Table(table_data, colWidths=[140, 150, 214])
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), primary_color),
                    ('TEXTCOLOR', (0,0), (-1,0), colors.white),
                    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#ced4da")),
                ]))
                story.append(t)
                story.append(Spacer(1, 8))
                in_table = False
                table_data = []

            num_text = md_to_reportlab_html(line_str)
            story.append(Paragraph(num_text, bullet_style))

        elif line_str.startswith("```") or line_str.startswith("+---") or (line_str.startswith("|") and len(line_str) > 1 and line_str[1] == '-'):
            story.append(Paragraph(html.escape(line_str).replace(" ", "&nbsp;"), code_style))

        elif line_str == "---":
            story.append(Spacer(1, 6))

        elif line_str != "":
            if in_table:
                t = Table(table_data)
                story.append(t)
                in_table = False
                table_data = []

            formatted_text = md_to_reportlab_html(line_str)
            story.append(Paragraph(formatted_text, body_style))

    doc.build(story)
    print(f"PDF generated successfully at {filename}")

def create_docx(filename):
    doc = Document()

    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("PLANTASYA MOBILE APPLICATION")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(27, 67, 50)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("OFFICIAL COMPREHENSIVE USER MANUAL & OPERATING GUIDE")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(45, 106, 79)

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = meta_p.add_run("App Name: Plantasya | Version: 1.0.0 | Target OS: Android 8.0+\nPackage: com.example.plantasya_mobileapp")
    run_meta.font.name = 'Calibri'
    run_meta.font.size = Pt(9.5)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(73, 80, 87)

    doc.add_paragraph("="*75)

    with open("PLANTASYA_USER_MANUAL.md", "r", encoding="utf-8") as f:
        lines = f.readlines()

    for line in lines:
        line_str = line.strip()
        if not line_str:
            continue

        if line_str.startswith("# "):
            h = doc.add_heading(md_to_docx_text(line_str[2:]), level=1)
            if len(h.runs) > 0:
                h.runs[0].font.color.rgb = RGBColor(27, 67, 50)
        elif line_str.startswith("## "):
            h = doc.add_heading(md_to_docx_text(line_str[3:]), level=2)
            if len(h.runs) > 0:
                h.runs[0].font.color.rgb = RGBColor(45, 106, 79)
        elif line_str.startswith("### "):
            h = doc.add_heading(md_to_docx_text(line_str[4:]), level=3)
            if len(h.runs) > 0:
                h.runs[0].font.color.rgb = RGBColor(64, 145, 108)
        elif line_str.startswith("* ") or line_str.startswith("- "):
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(md_to_docx_text(line_str[2:]))
        elif re.match(r'^\d+\.\s', line_str):
            p = doc.add_paragraph(style='List Number')
            p.add_run(md_to_docx_text(line_str))
        elif line_str == "---":
            continue
        else:
            p = doc.add_paragraph()
            p.add_run(md_to_docx_text(line_str))

    doc.save(filename)
    print(f"DOCX generated successfully at {filename}")

if __name__ == "__main__":
    create_pdf("PLANTASYA_USER_MANUAL.pdf")
    create_docx("PLANTASYA_USER_MANUAL.docx")
